from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path("legal/CAUSA_Legal_Terms_and_Privacy_Template.docx")


def shade(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def set_cell_text(cell, text: str, bold: bool = False) -> None:
    cell.text = ""
    run = cell.paragraphs[0].add_run(text)
    run.bold = bold
    run.font.size = Pt(9)


def paragraph(document, text: str) -> None:
    p = document.add_paragraph(text)
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.15


def heading(document, text: str, level: int) -> None:
    p = document.add_heading(text, level=level)
    p.paragraph_format.space_before = Pt(15 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(5)


def bullet(document, text: str) -> None:
    p = document.add_paragraph(style="List Bullet")
    p.add_run(text)
    p.paragraph_format.space_after = Pt(4)


def build() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = title.add_run("Cause AI Legal Terms and Privacy Template")
    run.font.name = "Arial"
    run.font.color.rgb = RGBColor(0, 0, 0)
    subtitle = doc.add_paragraph("Pre launch template for legal review and production completion")
    subtitle.runs[0].italic = True
    subtitle.paragraph_format.space_after = Pt(14)

    table = doc.add_table(rows=4, cols=2)
    table.style = "Table Grid"
    metadata = [
        ("Document owner", "[Company Legal Name and Legal Owner]"),
        ("Effective date", "[Effective Date after Counsel Approval]"),
        ("Version", "0.1 pre launch template"),
        ("Scope", "Cause AI financial exception investigation service"),
    ]
    for row, (label, value) in zip(table.rows, metadata):
        shade(row.cells[0], "17375E")
        set_cell_text(row.cells[0], label, True)
        for run in row.cells[0].paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
        set_cell_text(row.cells[1], value)
    paragraph(doc, "This template is not legal advice and does not guarantee legal compliance. It must be completed for the company’s actual legal entity, service model, jurisdictions, customer contracts, retention schedule, vendors and incident contacts, then approved by qualified counsel before publication.")

    heading(doc, "Completion Before Publication", 1)
    for item in [
        "Replace every bracketed placeholder and align the text with the signed customer agreement and data processing agreement.",
        "Confirm whether the company acts as data fiduciary, processor, payment aggregator, payment gateway, technology provider or another regulated role for each use case.",
        "Publish the approved privacy contact, grievance process, retention schedule, subprocessors, deployment regions and cross border transfer safeguards.",
        "Obtain privacy, security, regulatory, consumer protection and product counsel approval for every target jurisdiction.",
    ]:
        bullet(doc, item)

    heading(doc, "Terms of Use", 1)
    paragraph(doc, "These Terms govern authorised access to the Cause AI financial exception investigation service. The contracting entity is [Company Legal Name], with its registered address at [Registered Address]. The effective date is [Effective Date].")
    heading(doc, "Service Scope", 2)
    paragraph(doc, "Cause AI helps authorised business users identify payment-operation exceptions, assemble evidence and route investigations. It does not execute payments, hold customer funds, replace mandatory approvals or make a final fraud determination. Customers remain responsible for reconciliation, payment operations, legal compliance and final decisions.")
    heading(doc, "Accounts and Access", 2)
    paragraph(doc, "Customers must provide accurate account information, use least-privilege roles, safeguard credentials, promptly remove departed users and report suspected compromise. The provider may suspend access for a material security risk, misuse, legal requirement or material breach.")
    heading(doc, "Acceptable Use", 2)
    for item in [
        "Process only data the customer is authorised to provide and only for the approved business purpose.",
        "Do not upload PAN, CVV, full card-on-file data, third-party credentials, malware or data prohibited by law or contract.",
        "Do not bypass security controls, probe the service, interfere with availability or use outputs as the sole basis for a high-impact decision about an individual.",
    ]:
        bullet(doc, item)
    heading(doc, "Confidentiality and Intellectual Property", 2)
    paragraph(doc, "Each party must protect the other party’s confidential information with reasonable safeguards and use it only for the service relationship. Customers retain rights in customer data. The provider retains rights in the service, documentation and independently developed improvements.")
    heading(doc, "Commercial Terms", 2)
    paragraph(doc, "Service levels, support, fees, limitations of liability, warranties, indemnities, governing law, disputes and data processing must be stated in the signed customer agreement. This template does not create a banking, payment-service, legal, tax or regulatory advisory relationship.")

    heading(doc, "Privacy Notice", 1)
    paragraph(doc, "Data fiduciary and privacy contact: [Company Legal Name], [Registered Address], [Privacy Email]. Add the Data Protection Officer or grievance contact if required by applicable law.")
    heading(doc, "Personal Data Categories", 2)
    paragraph(doc, "The service may process business-user account details, contact information, tenant and role identifiers, payment references, amounts, currency, case evidence, audit events, device and security logs, and investigation comments. It must not store card PAN, CVV or full card-on-file data. Payment data must be minimised to what the stated investigation purpose requires.")
    heading(doc, "Purposes and Legal Basis", 2)
    paragraph(doc, "Data is processed to authenticate users, investigate exceptions, prevent fraud, maintain security, meet legal obligations, provide support and maintain audit records. The approved production notice must identify the applicable lawful basis, consent process where consent is relied upon, and purpose-specific notices.")
    heading(doc, "Sharing and Transfers", 2)
    paragraph(doc, "Data may be shared only with authorised personnel, contracted processors under written data-protection obligations, regulators or law-enforcement authorities where legally required, and parties approved by the customer. The final notice must list processor categories, storage regions and any cross-border transfer safeguards.")
    heading(doc, "Security Retention and Rights", 2)
    paragraph(doc, "Production must implement and publish a retention schedule, deletion process, legal-hold process, backup-expiry process and restoration test evidence. Individuals may request access, correction, erasure or grievance handling through [Privacy Email] or [Privacy Portal], subject to identity verification and applicable law. The production process must define response times, escalation and records of handling.")
    heading(doc, "Security Incidents", 2)
    paragraph(doc, "Suspected incidents must be assessed, contained, documented and reported to customers and authorities where required by law or contract. The production incident plan must name responsible contacts, decision authority, notification timelines and evidence-preservation steps.")

    heading(doc, "India Payment and Data Review", 1)
    paragraph(doc, "For India deployments, counsel and security owners must review applicability of the Digital Personal Data Protection Act 2023 and the phased Digital Personal Data Protection Rules 2025; RBI requirements applicable to the actual regulated role; restrictions on card-on-file storage; payment-system data localisation; CERT-In directions; and applicable NPCI or UPI obligations. Product statements must match implemented controls and regulatory status.")
    refs = doc.add_table(rows=1, cols=2)
    refs.style = "Table Grid"
    set_cell_text(refs.rows[0].cells[0], "Official source", True)
    set_cell_text(refs.rows[0].cells[1], "Review purpose", True)
    for cell in refs.rows[0].cells:
        shade(cell, "17375E")
        for run in cell.paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
    sources = [
        ("MeitY DPDP Act 2023 and DPDP Rules 2025", "Personal-data governance and phased compliance."),
        ("RBI payment-system data storage FAQ", "India storage and payment-data obligations."),
        ("RBI restriction on storage of actual card data", "Prevent prohibited card-data retention."),
    ]
    for source, purpose in sources:
        row = refs.add_row()
        set_cell_text(row.cells[0], source)
        set_cell_text(row.cells[1], purpose)
    paragraph(doc, "Template review date: 26 September 2026. Reconfirm current law and regulator guidance before every production release.")
    doc.save(OUTPUT)


if __name__ == "__main__":
    build()
