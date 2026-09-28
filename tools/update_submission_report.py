from docx import Document


REPORT = "reports/CAUSA_Production_Readiness_and_Legal_Submission_Report.docx"


def add_section(document, heading, paragraphs):
    document.add_heading(heading, level=1)
    for text in paragraphs:
        document.add_paragraph(text)


document = Document(REPORT)
marker = next(
    paragraph for paragraph in document.paragraphs
    if paragraph.text == "Verification Evidence"
)

sections = [
    ("Architecture Design", [
        "Cause AI uses a browser interface, Python HTTP API, workflow service, deterministic decision and risk engine, SQLite demonstration data store, append-only audit hash chain and regulatory-source monitor. Same-origin browser requests enter the API, which authenticates the user, applies role and tenant controls, then invokes the workflow service. The service returns only tenant-scoped cases, tickets, policy versions and audit events.",
        "The target production architecture replaces the local database and process state with managed PostgreSQL, company SSO, secret management, object storage for retained evidence, centralized logs and metrics, and a durable worker queue. These are release prerequisites and are not represented as deployed controls in the current demonstration."
    ]),
    ("Backend Design", [
        "The backend validates normalized payment, settlement, refund, fee, adjustment and bank records. It performs decimal financial calculations, creates immutable source-event hashes, prevents changed replayed records and stores cases under a tenant boundary. The detector attaches evidence identifiers, calculates explainable risk, applies the recorded policy version and stores the automatic decision with the case.",
        "The system records a separate evidence description for each decision. When escalation is required, it creates a department ticket containing an investigation description, risk factors, policy context, source evidence and linked prior case references. It does not execute payments, refunds, account changes or other financial actions."
    ]),
    ("Workflow Design", [
        "A validated source-record bundle is normalized and checked for exception patterns. The service creates a case with supplied User or Company ID, email address when present, evidence references and previous-report links. The deterministic engine then automatically approves, rejects or escalates the case. The case view exposes the decision rationale and the evidence description used to make it.",
        "Escalations create a department ticket automatically. The ticket tells the department what to investigate, why it was routed, which source records to reconcile and whether prior reports require review. A department records findings and a proposed resolution. An authorized analyst verifies a resolution against cited evidence. Material workflow actions are written to the append-only audit record."
    ]),
]

paragraphs = document.paragraphs
new_headings = {heading for heading, _ in sections}
start_index = min(i for i, paragraph in enumerate(paragraphs) if paragraph.text in new_headings)
end_index = next(i for i, paragraph in enumerate(paragraphs) if paragraph.text == "Verification Evidence")
new_section_paragraphs = paragraphs[start_index:end_index]

ordered = []
for heading, paragraphs in sections:
    ordered.append(("Heading 1", heading))
    ordered.extend(("Normal", text) for text in paragraphs)

if new_section_paragraphs:
    for paragraph, (style, text) in zip(new_section_paragraphs, ordered):
        paragraph.style = style
        paragraph.text = text
else:
    for heading, paragraphs in sections:
        heading_paragraph = document.add_heading(heading, level=1)
        marker._p.addprevious(heading_paragraph._p)
        for text in paragraphs:
            body = document.add_paragraph(text)
            marker._p.addprevious(body._p)

for paragraph in document.paragraphs:
    if paragraph.text.startswith("Automated test suite:"):
        paragraph.text = "Automated test suite: 74 passing tests. Checks include automatic case evaluation, evidence and investigation descriptions, access control, tenant isolation, detection decisions, append-only audit storage, immutable financial events, regulatory source safeguards, HTTP security behavior and Razorpay import failure handling."
    elif paragraph.text.startswith("Runtime evidence:"):
        paragraph.text = "Runtime evidence: local service health checks the current schema; the demo retains tenant-scoped cases and tickets; automatic evaluation is idempotent; the legal page and login-page legal link return HTTP 200; the Razorpay import endpoint returns HTTP 503 without company configuration. The production preflight evaluates twelve deployment prerequisites and fails closed when they are absent."

document.save(REPORT)
