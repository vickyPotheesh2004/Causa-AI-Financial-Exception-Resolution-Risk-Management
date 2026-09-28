from docx import Document

REPORT = "reports/CAUSA_Production_Readiness_and_Legal_Submission_Report.docx"
document = Document(REPORT)

replacements = {
    "Cause AI Production Readiness and Legal Submission Report": "Causa Submission Readiness and Demonstration Report",
    "Prepared for internal legal, privacy, security and business review": "Prepared for submission review and synthetic demonstration assessment",
    "Document date: 26 September 2026.": "Document date: 28 September 2026.",
    "Cross-document implementation score: 56 out of 100.": "Current submission status: SUBMISSION_READY_WITH_DOCUMENTED_LIMITATIONS.",
    "Automated test suite: 74 passing tests.": "Automated test suite: 76 passing tests.",
}
for paragraph in document.paragraphs:
    for old, new in replacements.items():
        if old in paragraph.text:
            paragraph.text = paragraph.text.replace(old, new)

sections = [
    ("Bounded AI Investigation", [
        "Causa provides an optional, zero-cost AI investigation layer. Local Ollama is the preferred provider. A Gemini adapter can run only when an account owner deliberately configures a free-tier-compatible model and key. When neither provider is available, Causa records a transparent deterministic fallback finding rather than claiming LLM output.",
        "The provider receives only a controlled case context: the exception, linked financial-event references, evidence, risk, policy, timeline and deterministic decision. It receives no database access, credentials or authority to change a financial outcome.",
        "AI output uses a strict schema for summary, root-cause status, cited evidence IDs, unknowns, conflicts, next steps and route recommendation. The backend validates cited evidence IDs against the supplied case context. AI findings remain advisory and cannot approve money movement, refund a payment, alter policy, change audit history, bypass verification or confirm fraud."
    ]),
    ("Zero Cost Demonstration Architecture", [
        "Local development runs with SQLite, local filesystem storage and optional Ollama. The public Vercel deployment runs synthetic data with ephemeral SQLite storage and deterministic fallback when no optional provider is configured. The deployment is suitable for demonstration only and must not receive real customer, payment or company data.",
        "A future company deployment needs durable PostgreSQL, persistent sessions, approved identity, object storage, secret management, monitoring, backups and company governance approval. These controls are not represented as deployed production controls in this report."
    ]),
    ("AI Safety Verification", [
        "The automated suite includes a prompt-injection test where an evidence record contains an instruction to approve a transaction. The system treats that text as untrusted record content and produces no financial authorization. A missing-evidence test confirms that deterministic fallback does not invent a bank record."
    ]),
]
paragraphs = document.paragraphs
existing_heads = {heading for heading, _ in sections}
start = min(i for i, p in enumerate(paragraphs) if p.text in existing_heads)
end = next(i for i, p in enumerate(paragraphs) if p.text == "Verification Evidence")
slots = paragraphs[start:end]
desired = []
for heading, contents in sections:
    desired.append(("Heading 1", heading))
    desired.extend(("Normal", content) for content in contents)
for paragraph, (style, text) in zip(slots, desired):
    paragraph.style = style
    paragraph.text = text

document.save(REPORT)
