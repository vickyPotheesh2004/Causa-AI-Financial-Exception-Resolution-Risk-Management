from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

OUT = Path("reports/CAUSA_Production_Readiness_and_Legal_Submission_Report.docx")

SECTIONS = [
    ("Executive Summary", [
        "Cause AI is a local financial-exception investigation application. It receives normalized finance records, detects discrepancies, creates tenant-scoped cases, routes investigations, records append-only audit events and supports controlled regulatory-source monitoring.",
        "The approved demonstration has strong workflow evidence, but it is not ready to process live Razorpay customer data or operate as a public production service. The current application deliberately blocks production launch and loopback-only enforcement prevents accidental exposure.",
    ]),
    ("Implemented Application Controls", [
        "The backend validates normalized record schema, uses decimal monetary calculations, blocks unsafe automated approval, routes suspicious activity to fraud review, links repeat requests to prior cases, and requires evidence before resolution verification.",
        "Tenant-scoped queries protect cases, tickets and audit events. Financial events are immutable, receive a SHA-256 content hash, have a unique source identity per tenant and reject replayed source records whose content changes.",
        "The HTTP layer enforces same-origin writes, content-security policy, no-store responses, MIME-sniffing protection, clickjacking protection, request-size limits, malformed Content-Length rejection, session expiry, session capacity limits and login throttling. Policy versions require separate proposal, review and activation actors, and each controlled tool declares roles, scope, rate limit and audit requirement.",
    ]),
    ("Razorpay and Financial Data Boundary", [
        "A read-only Razorpay Payments API connector exists. It requires explicit live mode, an authorised tenant match and a separate opt-in flag. It normalizes only payment fields needed for investigation and excludes raw email and card details from normalized records.",
        "The connector cannot run against live Razorpay data until an authorised company credential, tenant mapping, production secret store and operational approval are provided. The deployed endpoint fails closed with HTTP 503 when these controls are absent.",
    ]),
    ("Legal and Privacy Submission Package", [
        "The accompanying legal template covers service scope, authorised use, customer responsibilities, confidentiality, privacy categories, purposes, rights, sharing, retention, incidents and India-specific payment-data review. It contains required placeholders for the legal entity, contacts, processors, retention schedule, jurisdictions and contract terms.",
        "Counsel must complete and approve the template before publication. The document does not represent legal advice or a compliance guarantee.",
    ]),
    ("Production Readiness Assessment", [
        "Cross-document implementation score: 56 out of 100. This measures the current repository against the PDR and complete project description. Production readiness score: 55 out of 100. The score reflects verified application controls, policy governance and deployment preflight checks. It does not treat templates, planned architecture or an approved demonstration as deployed production controls.",
        "A score of 100 requires completed and evidenced company infrastructure, identity, approved data governance, independent assurance and authorised real integration testing.",
    ]),
    ("Release Blocking Gaps", [
        "Infrastructure: SQLite must be replaced by a managed production database with encryption, high availability, backup retention, restore testing and access logging. The service needs TLS ingress, WAF, private networking, secrets management, centralized logs, metrics and alerting.",
        "Identity and access: demo passwords and in-memory sessions must be replaced by a company identity provider with SSO, MFA, lifecycle provisioning, privileged-access controls and administrative audit trails.",
        "Financial integration: live Razorpay credentials, webhook signature verification, settlement, refund and dispute sources, approved reconciliation rules, idempotency monitoring and controlled production test evidence are still required.",
        "Security assurance: independent code review, threat model, VAPT, penetration testing, dependency scanning, load testing, disaster-recovery exercise, incident-response drill and operational ownership must be completed.",
        "Privacy and legal: signed terms, DPA, subprocessor register, retention and deletion configuration, privacy contact, data subject request workflow, data-localization evidence and counsel approval must be completed.",
    ]),
    ("Artificial Intelligence Use", [
        "The current application uses deterministic rules and explainable risk assessment rather than a generative model. It detects mismatches, duplicate records, timing exceptions and suspicious signals from normalized data. It preserves evidence identifiers and routes cases for human review.",
        "Any future ML model must undergo data-governance review, model-risk approval, evaluation against known outcomes, drift monitoring, human-override controls and documented adverse-impact assessment before it influences material decisions.",
    ]),
    ("Verification Evidence", [
        "Automated test suite: 72 passing tests. Checks include access control, tenant isolation, detection decisions, append-only audit storage, immutable financial events, regulatory source safeguards, HTTP security behavior and Razorpay import failure handling.",
        "Runtime evidence: local service health check reports schema version 4; the demo has nine retained cases after migration; the legal page and login-page legal link return HTTP 200; the Razorpay import endpoint returns HTTP 503 without company configuration. The production preflight evaluates twelve deployment prerequisites and fails closed when they are absent.",
    ]),
]

def p(doc, text):
    para = doc.add_paragraph(text)
    para.paragraph_format.space_after = Pt(7)
    para.paragraph_format.line_spacing = 1.15

def build():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(.7); sec.bottom_margin = Inches(.7); sec.left_margin = Inches(.8); sec.right_margin = Inches(.8)
    doc.styles['Normal'].font.name = 'Arial'; doc.styles['Normal'].font.size = Pt(10.5)
    t = doc.add_paragraph(style='Title'); t.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r=t.add_run('Cause AI Production Readiness and Legal Submission Report'); r.font.color.rgb=RGBColor(0,0,0)
    sub=doc.add_paragraph('Prepared for internal legal, privacy, security and business review'); sub.runs[0].italic=True
    p(doc, 'Document date: 26 September 2026. This report records implemented evidence and open release blockers. It does not provide legal advice, a security certification or a production approval.')
    for name, parts in SECTIONS:
        doc.add_heading(name, level=1)
        for part in parts: p(doc, part)
    doc.add_heading('Regulatory Reference Points', level=1)
    for source in [
        'MeitY: Digital Personal Data Protection Act 2023 and Digital Personal Data Protection Rules 2025.',
        'RBI: Restriction on Storage of Actual Card Data, RBI/2022-23/77.',
        'RBI: Storage of Payment System Data FAQ and relevant payment-aggregator directions.',
    ]: p(doc, source)
    doc.save(OUT)

if __name__ == '__main__': build()


