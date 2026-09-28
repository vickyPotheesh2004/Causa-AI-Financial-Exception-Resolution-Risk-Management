"""Add an evidence-based benchmark section to the existing submission report."""

from docx import Document


REPORT = "reports/CAUSA_Production_Readiness_and_Legal_Submission_Report.docx"


def insert_before(document: Document, marker, style: str, text: str) -> None:
    paragraph = document.add_paragraph(text, style=style)
    marker._p.addprevious(paragraph._p)


document = Document(REPORT)

replacements = {
    "The current application uses deterministic rules and explainable risk assessment rather than a generative model. It detects mismatches, duplicate records, timing exceptions and suspicious signals from normalized data. It preserves evidence identifiers and routes cases for human review.": (
        "Causa uses deterministic detection, risk assessment and decision controls, with an optional OpenRouter investigation provider. "
        "The provider is selected only when a server-side configuration and secret are present; otherwise the system records a transparent deterministic fallback finding. "
        "Generative output is advisory and cannot approve funds, override policy, modify an audit record or confirm fraud."
    ),
    "Causa provides an optional, zero-cost AI investigation layer. Local Ollama is the preferred provider. A Gemini adapter can run only when an account owner deliberately configures a free-tier-compatible model and key. When neither provider is available, Causa records a transparent deterministic fallback finding rather than claiming LLM output.": (
        "Causa provides an optional OpenRouter investigation layer. The API key is held only in server-side environment variables. "
        "When OpenRouter is not configured or the provider fails, Causa records a transparent deterministic fallback finding rather than claiming AI output."
    ),
    "Local development runs with SQLite, local filesystem storage and optional Ollama. The public Vercel deployment runs synthetic data with ephemeral SQLite storage and deterministic fallback when no optional provider is configured. The deployment is suitable for demonstration only and must not receive real customer, payment or company data.": (
        "Local development runs with SQLite, local filesystem storage and optional OpenRouter configuration. "
        "The public Vercel deployment runs synthetic data with ephemeral SQLite storage and uses the configured server-side provider or deterministic fallback. "
        "The deployment is suitable for demonstration only and must not receive real customer, payment or company data."
    ),
}

for paragraph in document.paragraphs:
    if paragraph.text in replacements:
        paragraph.text = replacements[paragraph.text]

# Keep this edit idempotent if the report is refreshed again.
if not any(paragraph.text == "Benchmark Status" for paragraph in document.paragraphs):
    marker = next(paragraph for paragraph in document.paragraphs if paragraph.text == "Verification Evidence")
    insert_before(document, marker, "Heading 1", "Benchmark Status")
    insert_before(
        document,
        marker,
        "Normal",
        "Measured on the included synthetic dataset: no benchmark result is available yet. The current repository does not contain the required 100-record labeled dataset, independent ground truth, or benchmark evaluator. Therefore precision, recall, F1, match rate, throughput and unresolved-case metrics have not been generated and are intentionally not reported.",
    )
    insert_before(
        document,
        marker,
        "Normal",
        "The benchmark acceptance plan is to evaluate deterministic exception detections against immutable synthetic ground truth by exception category. The evaluator must report total records, expected exceptions, true positives, false positives, false negatives, precision, recall, F1, match rate, unresolved case IDs and measured processing throughput. Results may be added to this report only after execution of that evaluator against the included dataset.",
    )
    insert_before(
        document,
        marker,
        "Normal",
        "Benchmark result: NOT MEASURED. This is a submission limitation and a release-blocking gap for any claim of measured detection quality.",
    )

document.save(REPORT)
